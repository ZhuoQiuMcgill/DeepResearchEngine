from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import math

TEMPLATE = '/mnt/d/Projects/DeepResearchEngine/docs/proposal/INSE 691 AI-Agent Sys. Project Proposal (Template).docx'

class Doc:
    def __init__(self, title, doc_kind, team, ids, date, body_pt=10.5, margin=0.8):
        self.doc = Document(TEMPLATE)
        body = self.doc.element.body
        children = list(body.iterchildren())
        date_idx = next(i for i, el in enumerate(children) if '[Pick the date]' in self._text(el))
        for el in children[date_idx + 1:]:
            if el.tag != qn('w:sectPr'): body.remove(el)
        for el in children[:date_idx + 1]:
            self._replace(el, 'Project Title', title)
            self._replace(el, 'Project Proposal', doc_kind)
            self._replace(el, 'Team Members', 'Team Members: ' + team)
            self._replace(el, 'Student IDs', 'Student IDs: ' + ids)
            self._replace(el, '[Pick the date]', date)
        sec = self.doc.sections[0]
        for side in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'): setattr(sec, side, Inches(margin))
        self.width = 8.5 - 2 * margin
        n = self.doc.styles['Normal']; n.font.size = Pt(body_pt); n.paragraph_format.space_after = Pt(5); n.paragraph_format.line_spacing = 1.05
        for hn, size, before, after in (('Heading 1', 13, 8, 2), ('Heading 2', 11.5, 6, 2)):
            st = self.doc.styles[hn]; st.font.size = Pt(size); st.paragraph_format.space_before = Pt(before); st.paragraph_format.space_after = Pt(after)
        self.body_pt = body_pt; self.est = []; self.bullet_num = self._bullet_numid()

    @staticmethod
    def _text(el): return ''.join(t.text or '' for t in el.iter(qn('w:t')))
    @staticmethod
    def _replace(el, old, new):
        for p in el.iter(qn('w:p')):
            ts = list(p.iter(qn('w:t'))); full = ''.join(t.text or '' for t in ts)
            if old in full and ts:
                ts[0].text = full.replace(old, new)
                for t in ts[1:]: t.text = ''
    def _bullet_numid(self):
        try: num_el = self.doc.part.numbering_part.element
        except Exception: return None
        abstract = {}
        for a in num_el.iter(qn('w:abstractNum')):
            lvl0 = next((l for l in a.iter(qn('w:lvl')) if l.get(qn('w:ilvl')) == '0'), None)
            fmt = lvl0.find(qn('w:numFmt')) if lvl0 is not None else None
            abstract[a.get(qn('w:abstractNumId'))] = fmt.get(qn('w:val')) if fmt is not None else None
        for n in num_el.iter(qn('w:num')):
            if abstract.get(n.find(qn('w:abstractNumId')).get(qn('w:val'))) == 'bullet': return n.get(qn('w:numId'))
        return None

    def h1(self, t): self.est.append(('h', t, self.width, 13)); return self.doc.add_paragraph(t, style='Heading 1')
    def h2(self, t): self.est.append(('h', t, self.width, 11.5)); return self.doc.add_paragraph(t, style='Heading 2')
    def h3(self, t):
        p = self.doc.add_paragraph(style='Normal'); r = p.add_run(t); r.bold = True; p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(2)
        self.est.append(('h', t, self.width, self.body_pt)); return p
    def para(self, t, italic=False, lead=None):
        p = self.doc.add_paragraph(style='Normal')
        if lead: r = p.add_run(lead); r.bold = True
        r = p.add_run(t); r.italic = italic
        self.est.append(('p', (lead or '') + t, self.width, self.body_pt)); return p
    def bullets(self, items):
        for it in items:
            if self.bullet_num:
                p = self.doc.add_paragraph(style='List Paragraph'); pPr = p._p.get_or_add_pPr(); numPr = OxmlElement('w:numPr')
                il = OxmlElement('w:ilvl'); il.set(qn('w:val'), '0'); ni = OxmlElement('w:numId'); ni.set(qn('w:val'), self.bullet_num)
                numPr.append(il); numPr.append(ni); pPr.append(numPr)
            else:
                p = self.doc.add_paragraph(style='Normal'); p.paragraph_format.left_indent = Inches(0.3); p.paragraph_format.first_line_indent = Inches(-0.2); p.add_run('•  ')
            if isinstance(it, tuple): r = p.add_run(it[0]); r.bold = True; p.add_run(it[1]); txt = it[0] + it[1]
            else: p.add_run(it); txt = it
            p.paragraph_format.space_after = Pt(2); self.est.append(('p', txt, self.width - 0.4, self.body_pt))
    @staticmethod
    def _borders(table):
        tblPr = table._tbl.tblPr; borders = OxmlElement('w:tblBorders')
        for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            e = OxmlElement(f'w:{edge}'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:space'), '0'); e.set(qn('w:color'), '999999'); borders.append(e)
        tblPr.append(borders)
    @staticmethod
    def _shade(cell, fill):
        tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill); tcPr.append(shd)
    @staticmethod
    def _cell(cell, txt, size, bold=False):
        cell.text = ''; p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
        r = p.add_run(txt); r.font.size = Pt(size); r.bold = bold
    def table(self, headers, rows, widths, font=9, first_bold=True):
        t = self.doc.add_table(rows=1, cols=len(headers)); self._borders(t)
        for i, h in enumerate(headers): self._cell(t.rows[0].cells[i], h, font, bold=True); self._shade(t.rows[0].cells[i], 'D9E2F3')
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row): self._cell(cells[i], str(v), font, bold=(first_bold and i == 0))
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Inches(w)
        sp = self.doc.add_paragraph(); sp.paragraph_format.space_after = Pt(0); sp.paragraph_format.line_spacing = 0.6
        self.est.append(('t', [headers] + rows, widths, font)); return t
    def gantt(self, acts, weeks, milestones, caption, label_w=2.6):
        g = self.doc.add_table(rows=1, cols=1 + weeks); self._borders(g)
        self._cell(g.rows[0].cells[0], 'Activity', 8, bold=True); self._shade(g.rows[0].cells[0], 'D9E2F3')
        for w in range(1, weeks + 1): self._cell(g.rows[0].cells[w], f'W{w}', 8, bold=True); self._shade(g.rows[0].cells[w], 'D9E2F3')
        for name, a, b in acts:
            cells = g.add_row().cells; self._cell(cells[0], name, 8)
            for w in range(1, weeks + 1):
                self._cell(cells[w], '', 8)
                if a <= w <= b: self._shade(cells[w], '4472C4')
        m = g.add_row().cells; self._cell(m[0], 'Milestones', 8, bold=True)
        for w in range(1, weeks + 1): self._cell(m[w], '', 8)
        for mid, w in milestones: self._cell(m[w], mid, 8, bold=True); self._shade(m[w], 'FFD966')
        cw = (self.width - label_w) / weeks
        for row in g.rows:
            row.cells[0].width = Inches(label_w)
            for w in range(1, weeks + 1): row.cells[w].width = Inches(cw)
        cap = self.doc.add_paragraph(style='Caption'); cap.add_run(caption)
        self.est.append(('t', [['x'] * (weeks + 1)] + [[n] for n, a, b in acts] + [['M']], [label_w] + [cw] * weeks, 8)); self.est.append(('p', caption, self.width, 9))
    def page_break(self):
        p = self.doc.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE); self.est.append(('pb', '', 0, 0))
    def references(self, refs):
        self.page_break(); self.h1('References')
        for i, r in enumerate(refs, 1):
            p = self.doc.add_paragraph(style='Normal'); p.paragraph_format.left_indent = Inches(0.4); p.paragraph_format.first_line_indent = Inches(-0.4); p.paragraph_format.space_after = Pt(3)
            run = p.add_run(f'[{i}] {r}'); run.font.size = Pt(9.5)
    def save(self, path): self.doc.save(path)
    def estimate(self, start_after_first_break=True):
        def lines_for(text, width_in, pt):
            cpl = max(10, int(width_in * 72 / (pt * 0.48))); return max(1, math.ceil(len(text) / cpl))
        total = 0.0; started = not start_after_first_break; breaks = 0
        for kind, txt, w, pt in self.est:
            if kind == 'pb':
                breaks += 1
                if breaks == 1 and start_after_first_break: started = True; continue
                if started: break
            if not started: continue
            lh = pt * 1.22 / 72
            if kind in ('p', 'h'): total += lines_for(txt, w, pt) * lh + (6 if kind == 'h' else 5) / 72
            elif kind == 't':
                for row in txt:
                    rl = max(lines_for(str(c), w[i] - 0.1, pt) for i, c in enumerate(row) if i < len(w)); total += rl * lh + 2 / 72
                total += 6 / 72
        usable = 11 - 2 * (8.5 - self.width) / 2
        return total, total / usable

REFS = [
    'Zeng, Y. (2004). Environment-based formulation of design problem. Journal of Integrated Design and Process Science, 8(4), 45-63. https://doi.org/10.3233/JID-2004-8404',
    'Zeng, Y. (2008). Recursive object model (ROM): Modelling of linguistic information in engineering design. Computers in Industry, 59(6), 612-625. https://doi.org/10.1016/j.compind.2008.03.002',
    'Wang, M., & Zeng, Y. (2009). Asking the right questions to elicit product requirements. International Journal of Computer Integrated Manufacturing, 22(4), 283-298. https://doi.org/10.1080/09511920802232902',
    'Zeng, Y. (2015). Environment-Based Design (EBD): A methodology for transdisciplinary design+. Journal of Integrated Design and Process Science, 19(1), 5-24. https://doi.org/10.3233/jid-2015-0004',
    'Mohammadi, A., & Zeng, Y. (2025). Enhancing TRIZ through environment-based design methodology supported by a large language model. AI EDAM, 39. https://doi.org/10.1017/S0890060425000083',
    'OpenAI. (2025). Introducing deep research. https://openai.com/index/introducing-deep-research/',
    'Google. (2025). Build with Gemini Deep Research. https://blog.google/technology/developers/deep-research-agent-gemini-api/',
    'Xu, R., & Peng, J. (2025). A comprehensive survey of deep research: Systems, methodologies, and applications. arXiv:2506.12594. https://arxiv.org/abs/2506.12594',
    'LangChain. (2025). Open Deep Research [source code]. https://github.com/langchain-ai/open_deep_research',
    'Chen, Z., Liu, K., Wang, Q., Liu, J., Zhang, W., Chen, K., & Zhao, F. (2024). MindSearch: Mimicking human minds elicits deep AI searcher. arXiv:2407.20183. https://arxiv.org/abs/2407.20183',
    'Du, M., Xu, B., Zhu, C., Wang, X., & Mao, Z. (2025). DeepResearch Bench: A comprehensive benchmark for deep research agents. arXiv:2506.11763. https://arxiv.org/abs/2506.11763',
    'Zheng, L., et al. (2023). Judging LLM-as-a-judge with MT-Bench and Chatbot Arena. Advances in Neural Information Processing Systems 36 (Datasets and Benchmarks). arXiv:2306.05685',
    'Wang, P., et al. (2023). Large language models are not fair evaluators. arXiv:2305.17926',
    'Hadfield, J., Zhang, B., Lien, K., Scholz, F., Fox, J., & Ford, D. (2025). How we built our multi-agent research system. Anthropic Engineering. https://www.anthropic.com/engineering/multi-agent-research-system',
]
