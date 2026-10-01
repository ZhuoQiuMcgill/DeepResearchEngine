from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import math

TEMPLATE = '/mnt/d/Projects/DeepResearchEngine/docs/INSE 691 AI-Agent Sys. Project Proposal (Template).docx'
OUT = '/mnt/d/Projects/DeepResearchEngine/docs/INSE691_Proposal_Draft_EBD-DeepResearch_v0.2.docx'
TITLE = 'EBD-DeepResearch: An Environment-Based Design Orchestrator for Multi-Agent Deep Research [proposed title]'

doc = Document(TEMPLATE)
body = doc.element.body
children = list(body.iterchildren())
def text_of(el): return ''.join(t.text or '' for t in el.iter(qn('w:t')))
date_idx = next(i for i, el in enumerate(children) if '[Pick the date]' in text_of(el))
for el in children[date_idx + 1:]:
    if el.tag != qn('w:sectPr'): body.remove(el)
def replace_in(el, old, new):
    for p in el.iter(qn('w:p')):
        ts = list(p.iter(qn('w:t'))); full = ''.join(t.text or '' for t in ts)
        if old in full and ts:
            ts[0].text = full.replace(old, new)
            for t in ts[1:]: t.text = ''
for el in children[:date_idx + 1]:
    replace_in(el, 'Project Title', TITLE)
    replace_in(el, 'Team Members', 'Team Members: [Member 1], [Member 2], [Member 3]')
    replace_in(el, 'Student IDs', 'Student IDs: [ID 1], [ID 2], [ID 3]')
    replace_in(el, '[Pick the date]', '[Date: to be set by the team]')

# page setup: slightly tighter margins and body type so the body fits three pages
sec = doc.sections[0]
for side in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
    setattr(sec, side, Inches(0.75))
normal = doc.styles['Normal']; normal.font.size = Pt(10)
normal.paragraph_format.space_after = Pt(4); normal.paragraph_format.line_spacing = 1.0
for hn, size, before, after in (('Heading 1', 12, 6, 1), ('Heading 2', 11, 4, 1)):
    st = doc.styles[hn]; st.font.size = Pt(size); st.paragraph_format.space_before = Pt(before); st.paragraph_format.space_after = Pt(after)
TEXT_WIDTH = 8.5 - 1.5

# ---------- helpers ----------
def bullet_numid():
    try: num_el = doc.part.numbering_part.element
    except Exception: return None
    abstract = {}
    for a in num_el.iter(qn('w:abstractNum')):
        lvl0 = next((l for l in a.iter(qn('w:lvl')) if l.get(qn('w:ilvl')) == '0'), None)
        fmt = lvl0.find(qn('w:numFmt')) if lvl0 is not None else None
        abstract[a.get(qn('w:abstractNumId'))] = fmt.get(qn('w:val')) if fmt is not None else None
    for n in num_el.iter(qn('w:num')):
        if abstract.get(n.find(qn('w:abstractNumId')).get(qn('w:val'))) == 'bullet': return n.get(qn('w:numId'))
    return None
BULLET_NUM = bullet_numid()
EST = []  # (kind, text, width_in, font_pt) for page estimate

def h1(t): EST.append(('h', t, TEXT_WIDTH, 13)); return doc.add_paragraph(t, style='Heading 1')
def h2(t): EST.append(('h', t, TEXT_WIDTH, 11.5)); return doc.add_paragraph(t, style='Heading 2')
def para(t, italic=False, bold_lead=None):
    p = doc.add_paragraph(style='Normal')
    if bold_lead: r = p.add_run(bold_lead); r.bold = True
    r = p.add_run(t); r.italic = italic
    EST.append(('p', (bold_lead or '') + t, TEXT_WIDTH, 10)); return p
def bullets(items):
    for it in items:
        if BULLET_NUM:
            p = doc.add_paragraph(style='List Paragraph'); pPr = p._p.get_or_add_pPr(); numPr = OxmlElement('w:numPr')
            il = OxmlElement('w:ilvl'); il.set(qn('w:val'), '0'); ni = OxmlElement('w:numId'); ni.set(qn('w:val'), BULLET_NUM)
            numPr.append(il); numPr.append(ni); pPr.append(numPr)
        else:
            p = doc.add_paragraph(style='Normal'); p.paragraph_format.left_indent = Inches(0.3); p.paragraph_format.first_line_indent = Inches(-0.2); p.add_run('•  ')
        if isinstance(it, tuple): r = p.add_run(it[0]); r.bold = True; p.add_run(it[1]); txt = it[0] + it[1]
        else: p.add_run(it); txt = it
        p.paragraph_format.space_after = Pt(2); EST.append(('p', txt, TEXT_WIDTH - 0.4, 10))
def set_borders(table):
    tblPr = table._tbl.tblPr; borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{edge}'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:space'), '0'); e.set(qn('w:color'), '999999'); borders.append(e)
    tblPr.append(borders)
def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill); tcPr.append(shd)
def cell_text(cell, txt, size, bold=False):
    cell.text = ''; p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
    r = p.add_run(txt); r.font.size = Pt(size); r.bold = bold
def table(headers, rows, widths, font=8.5, header_fill='D9E2F3', first_bold=False):
    t = doc.add_table(rows=1, cols=len(headers)); set_borders(t)
    for i, htxt in enumerate(headers): cell_text(t.rows[0].cells[i], htxt, font, bold=True); shade(t.rows[0].cells[i], header_fill)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row): cell_text(cells[i], str(v), font, bold=(first_bold and i == 0))
    for row in t.rows:
        for i, w in enumerate(widths): row.cells[i].width = Inches(w)
    sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(0); sp.paragraph_format.line_spacing = 0.6
    EST.append(('t', [headers] + rows, widths, font)); return t
def page_break():
    p = doc.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE); EST.append(('pb', '', 0, 0))

# ---------- cover note + page break ----------
para('Draft v0.2, 2026-09-17, for team review. Items in [square brackets] are to be decided by the team. The body is limited to three pages; the cover and references are excluded from the count.', italic=True)
page_break(); EST.clear()

h1('Objectives, Context and Scope')
h2('Objectives of the project')
para('EBD-DeepResearch is a multi-agent deep research system: given a research question it plans and runs web searches, reads sources, iterates and writes a cited report. Commercial products are closed [6, 7]; the open-source systems we surveyed (193 screened, 16 profiled) share one loop: free-form plan, parallel search, a sufficiency prompt, then a report. Their weaknesses are the failure modes named in Module 1: prose history serves as the database, stopping is a judgment call, and disagreement between sources is averaged away. Environment-Based Design (EBD) [1, 4] was built for this kind of formulation: analyse the environment, identify conflicts, generate a solution, repeat as the environment changes. Its Recursive Object Model [2] and question-asking strategy [3] define what must be known before a problem is understood; an LLM can carry out EBD activities [5]. No complete rule-based ROM parser exists, so we keep EBD\'s three units, objects, relations and environment, and let the LLM do the analysis.', bold_lead='Why. ')
bullets([
    ('O1. ', 'Build the system as an agent-based model: a coordinator and specialist seats with defined attributes, behaviours and interaction protocols, situated in a rule-managed environment, the ledger.'),
    ('O2. ', 'Implement an EBD coordinator and a matched baseline that share environment, tools, model adapter, budgets and writer, so any difference is attributable to EBD.'),
    ('O3. ', 'Evaluate both on DeepSeek-V4.1-Flash (API) and a local Qwen 3.8-27B (one RTX 5090) over nine tasks, judged blind by Claude Opus 5, GPT-5.6-sol and Gemini 3.8 Flash, with process metrics read from the ledger.'),
    ('O4. ', 'Deliver a Mission Console that replays any mission from the ledger, and documentation that lets the work grow into a paper.'),
])

h2('Context and Scope statement')
para('A chatbot answers fluently, but the user is the hidden orchestrator who decides every next question. A workflow fixes the path and breaks on the unusual case: a false presupposition in the question, or sources that contradict each other. An agent adds an evidence loop with state and stopping conditions. We split it for reasons the course accepts: context overload, since one fetched page can exceed a 128K to 256K window; independent parallel work per goal; and required independent checks. The coordinator keeps one point of control.', bold_lead='From chatbot to multi-agent. ')
para('Each seat is an agent with static attributes (role, instructions, tools, budgets, model) and dynamic attributes updated during the run. The seat map fixes who gives search goals, who calls the search tool and who decides what to fetch. The baseline uses the same seats and environment; its coordinator keeps prose notes instead of the task model and stops on a sufficiency check.', bold_lead='Agents (seats). ')
table(['Seat', 'Sees (sensors)', 'Decides and acts (behaviour)'], [
    ['Coordinator (EBD)', 'Task model, goal summaries; never raw pages', 'Analyses the task into objects, relations, environment; issues search goals; resolves conflicts; stops or escalates'],
    ['Researcher (per goal)', 'Its goal, result lists, its evidence records', 'Writes queries, calls the search tool, chooses what to fetch within rule limits, summarises the goal'],
    ['Distiller (rule + agent)', 'One goal, one page chunk', 'Rules estimate a page\'s token cost and split it to fit the budget; the agent extracts evidence spans with quotes'],
    ['Conflict Checker', 'Task model, new evidence', 'Classifies conflicts: source against source, refuted presupposition, constraint, scope, time; proposes resolutions'],
    ['Writer', 'Final task model, evidence', 'Writes the report, cites every claim, marks unresolved conflicts as disputed'],
], widths=[1.3, 1.9, 3.8], font=8, first_bold=True)
para('Agents are situated in a ledger (SQLite) run by fixed rules: software for exact work, the LLM for judgment. Rules store every result, page and evidence span with source and timestamp; deduplicate and cache; hold the task model and its status transitions; order goals, account for budgets and evaluate stopping conditions; validate agent outputs; and enforce a context gate, a token estimate before every agent call with splitting when a seat\'s budget would be exceeded. A seat\'s read and write rights on the ledger are its interaction with the environment. Every agent action and decision is appended as an event (seat, time, inputs, output, model, cost), so a mission can be audited and replayed.', bold_lead='Environment. ')
para('Two coordinators (baseline, EBD) by two models on nine tasks in three strata: relation-heavy; conflict-laden with planted contradictions and false presuppositions; open synthesis. Three runs per task [two if the schedule requires]. H1: at equal budget, EBD surfaces planted conflicts and false presuppositions more often. H2: blind judges prefer EBD reports on rubric coverage and citation support, or find them no worse. H3: EBD stops by its own condition more often and needs fewer searches per useful source. H4: the EBD gain is larger on the local model.', bold_lead='Comparison and hypotheses. ')
para('In: both coordinators, the seats, the ledger and tool layer (open-source search, fetch, cache), the Mission Console, evaluation, tests, documentation. Out: model training, browser automation, paywalled sources, a chat product, non-English or multimodal tasks.', bold_lead='Scope. ')

h2('Agent-based modelling questions (Module 1, slides 23 to 25)')
table(['Question', 'Answer for EBD-DeepResearch'], [
    ['1. Problem, questions, value added', 'Cited research where evidence is dispersed, conflicting, or contradicts the question\'s assumptions; the model must say what is presupposed, unknown, supported, conflicting, and when research is done. Interacting agents let each search depend on the last result and resolve disagreement, which a chatbot (user-orchestrated) and a workflow (fixed path) cannot.'],
    ['2. Agents, decision makers, attributes', 'The five seats. Decisions: Coordinator (goals, stopping), Researchers (queries, fetches), Distiller (what is evidence), Checker (what conflicts), Writer (structure). Static: role, instructions, tools, budgets, model. Dynamic: task model statuses, goal list, conflicts, evidence, budget spent.'],
    ['3. Environment and interaction', 'External: the web through search and fetch tools, plus the asker\'s purpose, time frame and scope. Shared: the ledger, also the EBD environment that grows each round. Seats sense it through read rights and act through write rights. Mobility is movement through information space, bounded by a goal.'],
    ['4. Behaviours, decisions, actions', 'Coordinator: analyse, prioritise, dispatch, resolve, stop. Researcher: query, select, fetch, summarise. Distiller: extract. Checker: compare, classify. Writer: synthesise, cite. Exact work is rule-based (budgets, ordering, chunking, checks, stopping); judgment is LLM-based; the goal list adapts to evidence.'],
    ['5. Interactions', 'Hub and spoke through the Coordinator: goals down, summaries up, evidence to the Checker, conflicts back, final model to the Writer. Researchers never talk to each other; interdependency runs through the ledger; the outside world is reached only through tools.'],
    ['6. Data sources', 'Behaviours: prompts and patterns from the profiled open-source systems (baseline), the EBD literature (coordinator), and our own ledger traces. Environment: the live web, frozen per run by the cache. Tasks: nine authored tasks with rubrics plus development tasks.'],
    ['7. Validation', 'Behaviour level: unit tests of rules; task analysis checked against hand-made decompositions and for stability; Checker scored on planted conflicts. System level: the two-by-two comparison with blind judges, rubric coverage, citation support, conflict recall, self-termination, cost, paired statistics and a human check. Trace level: every run replays from the ledger and answers what was observed, calculated, by which rule, from which source and time.'],
], widths=[1.35, 5.65], font=8, first_bold=True)

h2('Stakeholder profiles')
para('End users (graduate students, analysts, engineers) who need cited, traceable reports that state disagreements. The instructor and teaching assistant, who expect a working multi-agent system with testing and controlled execution. Prof. Yong Zeng\'s design lab, owner of EBD and potential co-author, expecting evidence that EBD improves information seeking. The team as first users, and the open-source community reusing the released tasks, prompts and code.')

h1('Product and deliverables')
bullets([
    ('D1. System (open license). ', 'Command line and Python API; model adapter for both endpoints; the five seats; the ledger with tool layer, context gate and event log; both coordinators.'),
    ('D2. Mission Console. ', 'Replays a mission from the ledger: a timeline of every event (agent, time, tool call, query and topic, pages, cost), the evolving task model and conflicts, two missions side by side.'),
    ('D3. Evaluation kit. ', 'Nine tasks with rubrics, judge prompts, scripts for blind pairwise judging, rubric scoring, trace metrics and statistics; the results dataset.'),
    ('D4. Testing report. ', 'Unit tests of the rule layer, end-to-end integration tests in both modes on both models, evaluation results with inter-judge agreement.'),
    ('D5. Documentation. ', 'User manual, developer guide, final report and presentation, written to extend into a paper.'),
])

h1('Important tasks to be completed')
para('T1 literature and design (seat map, protocols, ledger schema); T2 ledger and tool layer (search, fetch, cache, context gate, event log); T3 model adapter and local serving; T4 baseline coordinator; T5 EBD coordinator; T6 shared seats (Researcher, Distiller, Checker, Writer); T7 Mission Console; T8 tasks, rubrics, judges, metrics; T9 runs on both models and statistics; T10 testing, documentation, report. Owners, effort and timing: Figure 1.')

h1('Project process activities, schedule and milestones')
para('Weekly increments, each ending with a working system and a demonstration; two integration points and a scope freeze. Weeks count from the project start [Week 1: TBD; final report due mid-December 2026]. Milestones: M1 proposal approved (W1); M2 ledger, tools and baseline end to end (W3); M3 EBD end to end on both models (W5); M4 tasks and rubrics frozen, scope freeze (W6); M5 all runs judged (W8); M6 report and presentation (W9). Effort in person-days.', bold_lead='Process. ')
ACTS = [('T1', 'Literature and design', 'All', 1, 2, 6), ('T2', 'Ledger and tool layer', 'M1', 1, 3, 8), ('T3', 'Model adapter, local serving', 'M1', 2, 3, 4),
        ('T4', 'Baseline coordinator', 'M1', 2, 3, 5), ('T5', 'EBD coordinator', 'M2', 2, 5, 10), ('T6', 'Shared seats', 'M2', 3, 5, 8),
        ('T7', 'Mission Console', 'M1+all', 4, 8, 9), ('T8', 'Tasks, rubrics, judges, metrics', 'M3', 2, 6, 10), ('T9', 'Runs and statistics', 'M3+all', 6, 8, 8),
        ('T10', 'Testing, docs, report', 'All', 5, 9, 12)]
WEEKS = 9
g = doc.add_table(rows=1, cols=4 + WEEKS); set_borders(g)
for i, txt in enumerate(['ID', 'Activity', 'Owner', 'Days'] + [f'W{w}' for w in range(1, WEEKS + 1)]):
    cell_text(g.rows[0].cells[i], txt, 8, bold=True); shade(g.rows[0].cells[i], 'D9E2F3')
for a in ACTS:
    cells = g.add_row().cells
    cell_text(cells[0], a[0], 8, bold=True); cell_text(cells[1], a[1], 8); cell_text(cells[2], a[2], 8); cell_text(cells[3], str(a[5]), 8)
    for w in range(1, WEEKS + 1):
        cell_text(cells[3 + w], '', 8)
        if a[3] <= w <= a[4]: shade(cells[3 + w], '4472C4')
MS = [('M1', 1), ('M2', 3), ('M3', 5), ('M4', 6), ('M5', 8), ('M6', 9)]
mrow = g.add_row().cells
cell_text(mrow[0], 'M', 8, bold=True); cell_text(mrow[1], 'Milestones', 8); cell_text(mrow[2], '', 8); cell_text(mrow[3], '', 8)
for w in range(1, WEEKS + 1): cell_text(mrow[3 + w], '', 8)
for m, w in MS: cell_text(mrow[3 + w], m, 8, bold=True); shade(mrow[3 + w], 'FFD966')
for row in g.rows:
    for i, wd in enumerate([0.4, 2.0, 0.6, 0.45] + [0.4] * WEEKS): row.cells[i].width = Inches(wd)
cap = doc.add_paragraph(style='Caption'); cap.add_run('Figure 1. Gantt chart by project week: shaded cells are planned activity, M1 to M6 are milestones.')
EST.append(('t', [['x'] * 13] + [[a[0], a[1], a[2], '', ''] for a in ACTS] + [['M']], [0.4, 2.0, 0.6, 0.45] + [0.4] * WEEKS, 8)); EST.append(('p', 'caption', TEXT_WIDTH, 9))

h1('Project team organization')
table(['Member', 'Role', 'Responsibilities'], [
    ['[Member 1]', 'Platform lead', 'Ledger, tool layer, context gate, model adapter and local serving, baseline coordinator, Mission Console backend, continuous integration'],
    ['[Member 2]', 'EBD lead', 'EBD coordinator, seat prompts and behaviours, conflict typology, Writer integration'],
    ['[Member 3]', 'Evaluation lead', 'Task set and rubrics, judge pipeline, trace metrics, statistics, testing report'],
], widths=[1.0, 1.2, 4.8], font=8, first_bold=True)
para('Shared: console front end [owner TBD], weekly integration, report and presentation. Area leads decide within their area; cross-cutting decisions by majority in the weekly meeting; the instructor is consulted at milestones.')

h1('Development risks')
table(['ID', 'Risk', 'L/I', 'Mitigation'], [
    ['R1', 'Local Qwen weaker at structured output and tool use than the API model', 'H/H', 'Schema validation and repair in the rule layer; short structured prompts; API model as reference; report the gap (H4)'],
    ['R2', 'Local throughput (about 100 tokens per second) bounds run time', 'H/M', 'Context and fetch caps per seat; overnight runs; fall back to two runs per task'],
    ['R3', 'Search unreliability: engine blocked or API limits', 'M/H', 'Cache every result; paid search API as fallback; fixed experiment window'],
    ['R4', 'Judge bias (position, length, family) and disagreement', 'H/H', 'Blind judging, position swap, rubrics, agreement statistics, small human check; judges from other model families'],
    ['R5', 'Nine tasks limit power; task analysis varies between runs', 'H/M', 'Repeated runs, paired tests, pilot-scale claims; fixed decoding and validation'],
    ['R6', 'Time (nine weeks, three people) and console scope creep', 'H/H', 'Scope freeze at M4; minimum console first; extras only after M5'],
], widths=[0.4, 2.6, 0.4, 3.6], font=8, first_bold=True)

h1('Control methods')
bullets([
    'Weekly meeting with milestone check, task-board burndown, written minutes and a risk review.',
    'Git with feature branches and pull-request review; continuous integration runs unit and integration tests on every change.',
    'The ledger is the experiment log: every mission is recorded with configuration, events and cost; results are frozen once judged.',
    'Pre-registered hypotheses and metrics (this proposal); scope freeze at M4; instructor checkpoints at M2, M4 and M6.',
])

# ---------- References (own page) ----------
page_break()
h1('References')
REFS = [
    'Zeng, Y. (2004). Environment-based formulation of design problem. Journal of Integrated Design and Process Science, 8(4), 45-63. https://doi.org/10.3233/JID-2004-8404',
    'Zeng, Y. (2008). Recursive object model (ROM): Modelling of linguistic information in engineering design. Computers in Industry, 59(6), 612-625. https://doi.org/10.1016/j.compind.2008.03.002',
    'Wang, M., & Zeng, Y. (2009). Asking the right questions to elicit product requirements. International Journal of Computer Integrated Manufacturing, 22(4), 283-298. https://doi.org/10.1080/09511920802232902',
    'Zeng, Y. (2015). Environment-Based Design (EBD): A methodology for transdisciplinary design+. Journal of Integrated Design and Process Science, 19(1), 5-24. https://doi.org/10.3233/jid-2015-0004',
    'Mohammadi, A., & Zeng, Y. (2025). Enhancing TRIZ through environment-based design methodology supported by a large language model. AI EDAM, 39. https://doi.org/10.1017/S0890060425000083',
    'OpenAI. (2025). Introducing deep research. https://openai.com/index/introducing-deep-research/',
    'Google. (2025). Build with Gemini Deep Research. https://blog.google/technology/developers/deep-research-agent-gemini-api/',
    'Xu, R., & Peng, J. (2025). A comprehensive survey of deep research: Systems, methodologies, and applications. arXiv:2506.12594. https://arxiv.org/abs/2506.12594',
    'LangChain. (2025). Open Deep Research [source code]. https://github.com/langchain-ai/open_deep_research',
    'Chen, Z., Liu, K., Wang, Q., Liu, J., Zhang, W., Chen, K., & Zhao, F. (2024). MindSearch: Mimicking human minds elicits deep AI searcher. arXiv:2407.20183. https://arxiv.org/abs/2407.20183',
    'Du, M., Xu, B., Zhu, C., Wang, X., & Mao, Z. (2025). DeepResearch Bench: A comprehensive benchmark for deep research agents. arXiv:2506.11763. https://arxiv.org/abs/2506.11763',
    'Zheng, L., et al. (2023). Judging LLM-as-a-judge with MT-Bench and Chatbot Arena. Advances in Neural Information Processing Systems 36 (Datasets and Benchmarks). arXiv:2306.05685',
    'Wang, P., et al. (2023). Large language models are not fair evaluators. arXiv:2305.17926',
    'Hadfield, J., Zhang, B., Lien, K., Scholz, F., Fox, J., & Ford, D. (2025). How we built our multi-agent research system. Anthropic Engineering. https://www.anthropic.com/engineering/multi-agent-research-system',
]
for i, r in enumerate(REFS, 1):
    p = doc.add_paragraph(style='Normal'); p.paragraph_format.left_indent = Inches(0.4); p.paragraph_format.first_line_indent = Inches(-0.4); p.paragraph_format.space_after = Pt(3)
    run = p.add_run(f'[{i}] {r}'); run.font.size = Pt(9.5)
doc.save(OUT)

# ---------- rough page estimate for the body (between the two page breaks) ----------
def lines_for(text, width_in, pt):
    cpl = max(10, int(width_in * 72 / (pt * 0.5)))  # ~0.5em average glyph width
    return max(1, math.ceil(len(text) / cpl))
total_in = 0.0
for kind, txt, w, pt in EST:
    if kind == 'pb': break
    lh = pt * 1.25 / 72
    if kind in ('p', 'h'): total_in += lines_for(txt, w, pt) * lh + (6 if kind == 'h' else 4) / 72
    elif kind == 't':
        for row in txt:
            rl = max(lines_for(str(c), w[i] - 0.1, pt) for i, c in enumerate(row) if i < len(w))
            total_in += rl * lh + 3 / 72
        total_in += 6 / 72
usable = 11 - 1.5
print(f'estimated body height: {total_in:.1f} in  ->  about {total_in / usable:.2f} pages (usable {usable:.1f} in/page)')
